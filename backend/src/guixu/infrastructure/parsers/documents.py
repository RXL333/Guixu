from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from guixu.domain.profiles import Coverage, Evidence, EvidenceLocator, FileProfile, ParseOutcome
from guixu.infrastructure.parsers.common import PARSER_VERSION, ParseLimitError, stratified_indices, validate_zip_container
from guixu.infrastructure.parsers.ocr import recognize


PRESET_LIMITS = {"fast": {"pdf": 4, "pptx": 6, "xlsx": 3}, "standard": {"pdf": 12, "pptx": 20, "xlsx": 8}, "deep": {"pdf": 30, "pptx": 60, "xlsx": 20}}


def parse_pdf(path: Path, file_id: str, preset: str) -> ParseOutcome:
    warnings: list[str] = []
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        profile = FileProfile(file_id=file_id, modality="document", document_kind="pdf", metadata={"encrypted": True}, coverage=Coverage(mode="metadata_only", total_pages=None), warnings=["PDF_ENCRYPTED"], capabilities_used=["pypdf"], parser_version=PARSER_VERSION)
        return ParseOutcome(status="partial", profile=profile)
    total = len(reader.pages)
    selected = stratified_indices(total, PRESET_LIMITS[preset]["pdf"])
    evidence = []
    low_text_pages = []
    for index in selected:
        text = (reader.pages[index].extract_text() or "").strip()
        if len(text) < 20:
            low_text_pages.append(index + 1)
            continue
        evidence.append(Evidence(id=f"page-{index + 1}", kind="extracted_text", text=text[:12_000], locator=EvidenceLocator(page=index + 1), quality="high", origin="pypdf"))
    if low_text_pages:
        try:
            import pypdfium2 as pdfium
            pdf_document = pdfium.PdfDocument(path)
            for page_number in low_text_pages:
                rendered = pdf_document[page_number - 1].render(scale=2).to_pil().convert("RGB")
                lines, warning = recognize(rendered)
                if lines:
                    evidence.append(Evidence(id=f"page-{page_number}-ocr", kind="ocr", text="\n".join(line for line, _ in lines)[:12_000], locator=EvidenceLocator(page=page_number), quality="high" if min(score for _, score in lines) >= 0.8 else "medium", origin="RapidOCR-local-verified"))
                elif warning:
                    warnings.append(warning)
                else:
                    warnings.append(f"PDF_OCR_NO_TEXT:{page_number}")
        except Exception:
            warnings.append("PDF_RENDER_OR_OCR_FAILED")
    truncated = len(selected) < total
    profile = FileProfile(
        file_id=file_id, modality="document", document_kind="pdf", metadata={"encrypted": False, "pages": total},
        content_summary="\n".join(item.text for item in evidence)[:2000], summary_origin="deterministic" if evidence else "none",
        evidence=evidence, coverage=Coverage(mode="sampled" if truncated or low_text_pages else "full", total_pages=total, sampled_pages=[index + 1 for index in selected], truncated=truncated),
        warnings=warnings, capabilities_used=["pypdf"] + (["pypdfium2", "RapidOCR-CPU"] if any(item.kind == "ocr" for item in evidence) else []), parser_version=PARSER_VERSION,
    )
    return ParseOutcome(status="partial" if warnings else "ready", profile=profile)


def parse_docx(path: Path, file_id: str, preset: str) -> ParseOutcome:
    del preset
    validate_zip_container(path)
    document = Document(path)
    evidence = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            evidence.append(Evidence(id=f"paragraph-{index}", kind="extracted_text", text=text[:12_000], locator=EvidenceLocator(paragraph=index), quality="high", origin="python-docx"))
    for table_index, table in enumerate(document.tables, 1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows[:50]]
        if rows:
            evidence.append(Evidence(id=f"table-{table_index}", kind="extracted_text", text="\n".join(rows)[:12_000], locator=EvidenceLocator(field=f"table:{table_index}"), quality="high", origin="python-docx"))
    profile = FileProfile(file_id=file_id, modality="document", document_kind="docx", metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)}, content_summary="\n".join(item.text for item in evidence)[:2000], summary_origin="deterministic" if evidence else "none", evidence=evidence, coverage=Coverage(mode="full"), capabilities_used=["python-docx"], parser_version=PARSER_VERSION)
    return ParseOutcome(status="ready", profile=profile)


def parse_pptx(path: Path, file_id: str, preset: str) -> ParseOutcome:
    validate_zip_container(path)
    presentation = Presentation(path)
    total = len(presentation.slides); selected = stratified_indices(total, PRESET_LIMITS[preset]["pptx"])
    evidence = []
    for index in selected:
        texts = []
        for shape in presentation.slides[index].shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                texts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows[:30])
        if texts:
            evidence.append(Evidence(id=f"slide-{index + 1}", kind="extracted_text", text="\n".join(texts)[:12_000], locator=EvidenceLocator(slide=index + 1), quality="high", origin="python-pptx"))
    truncated = len(selected) < total
    profile = FileProfile(file_id=file_id, modality="document", document_kind="pptx", metadata={"slides": total}, content_summary="\n".join(item.text for item in evidence)[:2000], summary_origin="deterministic" if evidence else "none", evidence=evidence, coverage=Coverage(mode="sampled" if truncated else "full", total_pages=total, sampled_pages=[i + 1 for i in selected], truncated=truncated), warnings=["PPTX_SLIDES_SAMPLED"] if truncated else [], capabilities_used=["python-pptx"], parser_version=PARSER_VERSION)
    return ParseOutcome(status="partial" if truncated else "ready", profile=profile)


def parse_xlsx(path: Path, file_id: str, preset: str) -> ParseOutcome:
    validate_zip_container(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheet_names = list(workbook.sheetnames)
    selected_sheets = sheet_names[: PRESET_LIMITS[preset]["xlsx"]]
    evidence = []; warnings = []; formula_count = 0
    row_limit = {"fast": 20, "standard": 50, "deep": 200}[preset]
    for sheet_name in selected_sheets:
        sheet = workbook[sheet_name]
        sampled = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=False), 1):
            if row_index > row_limit:
                warnings.append(f"XLSX_ROWS_SAMPLED:{sheet_name}")
                break
            values = []
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                    values.append("[FORMULA_NOT_EXECUTED]")
                elif value is not None:
                    values.append(str(value))
                else:
                    values.append("")
            if any(values):
                sampled.append((row_index, " | ".join(values)))
        if sampled:
            rows = [item[0] for item in sampled]
            evidence.append(Evidence(id=f"sheet-{len(evidence)+1}", kind="extracted_text", text="\n".join(item[1] for item in sampled)[:12_000], locator=EvidenceLocator(sheet=sheet_name, rows=[rows[0], rows[-1]]), quality="high", origin="openpyxl"))
    workbook.close()
    if formula_count:
        warnings.append(f"XLSX_FORMULAS_NOT_EXECUTED:{formula_count}")
    truncated = len(selected_sheets) < len(sheet_names) or any(item.startswith("XLSX_ROWS_SAMPLED") for item in warnings)
    profile = FileProfile(file_id=file_id, modality="document", document_kind="xlsx", metadata={"sheets": sheet_names, "formula_count": formula_count}, content_summary="\n".join(item.text for item in evidence)[:2000], summary_origin="deterministic" if evidence else "none", evidence=evidence, coverage=Coverage(mode="sampled" if truncated else "full", truncated=truncated), warnings=warnings, capabilities_used=["openpyxl-read-only"], parser_version=PARSER_VERSION)
    return ParseOutcome(status="partial" if warnings else "ready", profile=profile)


DOCUMENT_PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".pptx": parse_pptx, ".xlsx": parse_xlsx}
