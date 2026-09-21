from __future__ import annotations

import csv
import json
import math
import hashlib
import os
import sys
import time
import struct
import uuid
import wave
import zipfile
from pathlib import Path

import pytest
import cv2
import numpy as np
from docx import Document
from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from guixu.application.parser_runner import ParserRunner
from guixu.infrastructure.db.database import Database
from guixu.infrastructure.parsers.registry import ParserRegistry
from guixu.infrastructure.resources.components import ComponentError, ComponentManager


def uid() -> str:
    return str(uuid.uuid4())


def assert_schema(project_root: Path, outcome) -> None:
    schema = json.loads((project_root / "contracts" / "schemas" / "file-profile.schema.json").read_text("utf-8"))
    Draft202012Validator(schema, format_checker=FormatChecker()).validate(outcome.profile.model_dump(mode="json"))
    assert outcome.profile.name
    assert outcome.profile.source_path
    assert outcome.profile.extension.startswith(".")
    assert outcome.profile.parser_status == outcome.status
    assert outcome.profile.parser_warnings == outcome.profile.warnings


def test_pa01_text_markdown_json_csv(project_root: Path, tmp_path: Path):
    fixtures = {
        "中文.txt": "标题\n\n第一段\n\n最后一段",
        "notes.md": "# Heading\n\nDo not execute <script>alert(1)</script>",
        "data.json": json.dumps({"kind": "invoice", "items": [1, 2]}, ensure_ascii=False),
        "table.csv": "name,value\nalpha,1\nbeta,2\n",
    }
    for name, content in fixtures.items():
        path = tmp_path / name; path.write_text(content, encoding="utf-8")
        outcome = ParserRegistry().parse(path, uid())
        assert outcome.status == "ready" and outcome.profile.evidence
        assert_schema(project_root, outcome)


def test_pa02_text_pdf_scanned_pdf_and_encrypted_pdf(project_root: Path, tmp_path: Path):
    text_pdf = tmp_path / "text.pdf"
    pdf = canvas.Canvas(str(text_pdf)); pdf.drawString(72, 720, "Guixu parser evidence on a real PDF page."); pdf.save()
    text_outcome = ParserRegistry().parse(text_pdf, uid())
    assert text_outcome.status == "ready" and text_outcome.profile.evidence[0].locator.page == 1
    scanned_pdf = tmp_path / "scan.pdf"
    scan_image = Image.new("RGB", (1000, 300), "white")
    ImageDraw.Draw(scan_image).text((40, 90), "SCANNED INVOICE 2026", fill="black", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 64))
    scan_image.save(scanned_pdf, "PDF")
    scan_outcome = ParserRegistry().parse(scanned_pdf, uid())
    assert any(item.kind == "ocr" and "INVOICE" in item.text for item in scan_outcome.profile.evidence)
    assert "RapidOCR-CPU" in scan_outcome.profile.capabilities_used
    encrypted = tmp_path / "encrypted.pdf"; writer = PdfWriter(); writer.add_blank_page(100, 100); writer.encrypt("secret")
    with encrypted.open("wb") as handle: writer.write(handle)
    encrypted_outcome = ParserRegistry().parse(encrypted, uid())
    assert encrypted_outcome.status == "partial" and encrypted_outcome.profile.warnings == ["PDF_ENCRYPTED"]
    for outcome in (text_outcome, scan_outcome, encrypted_outcome): assert_schema(project_root, outcome)


def test_pa03_docx_table_and_pptx_titles(project_root: Path, tmp_path: Path):
    docx_path = tmp_path / "document.docx"; document = Document(); document.add_heading("Project Atlas", 1)
    table = document.add_table(rows=2, cols=2); table.cell(0, 0).text = "Owner"; table.cell(0, 1).text = "Status"; table.cell(1, 0).text = "Lin"; table.cell(1, 1).text = "Ready"; document.save(docx_path)
    docx_outcome = ParserRegistry().parse(docx_path, uid()); assert docx_outcome.status == "ready"; assert any("Owner" in item.text for item in docx_outcome.profile.evidence)
    pptx_path = tmp_path / "slides.pptx"; presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Quarterly Review"; slide.placeholders[1].text = "Evidence body"; presentation.save(pptx_path)
    pptx_outcome = ParserRegistry().parse(pptx_path, uid()); assert pptx_outcome.status == "ready"; assert "Quarterly Review" in pptx_outcome.profile.content_summary
    assert_schema(project_root, docx_outcome); assert_schema(project_root, pptx_outcome)


def test_pa04_xlsx_does_not_execute_formula(project_root: Path, tmp_path: Path):
    path = tmp_path / "sheet.xlsx"; workbook = Workbook(); sheet = workbook.active; sheet.title = "Budget"; sheet.append(["Item", "Amount"]); sheet.append(["Total", "=SUM(B3:B4)"]); workbook.save(path)
    outcome = ParserRegistry().parse(path, uid())
    assert outcome.status == "partial"
    assert outcome.profile.metadata["formula_count"] == 1
    assert "[FORMULA_NOT_EXECUTED]" in outcome.profile.evidence[0].text
    assert_schema(project_root, outcome)


def test_pa05_image_thumbnail_strips_exif_and_pa06_gif_samples_frames(project_root: Path, tmp_path: Path):
    image_path = tmp_path / "photo.jpg"; exif = Image.Exif(); exif[271] = "Test Camera"; exif[34853] = {1: "sensitive"}
    source_image = Image.new("RGB", (900, 300), "white")
    ImageDraw.Draw(source_image).text((30, 90), "GUIXU RECEIPT 2026", fill="black", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 58))
    source_image.save(image_path, exif=exif)
    image_outcome = ParserRegistry().parse(image_path, uid(), artifact_dir=tmp_path / "thumbs")
    thumbnail = Path(image_outcome.cache_artifacts[0]); assert thumbnail.exists()
    with Image.open(thumbnail) as generated: assert not generated.getexif()
    assert "gps" not in json.dumps(image_outcome.profile.metadata).lower()
    assert any(item.kind == "ocr" and "GUIXU" in item.text for item in image_outcome.profile.evidence)
    gif_path = tmp_path / "animated.gif"; frames = [Image.new("RGB", (40, 40), (index * 40, 0, 0)) for index in range(5)]
    frames[0].save(gif_path, save_all=True, append_images=frames[1:], duration=100, loop=0)
    gif_outcome = ParserRegistry().parse(gif_path, uid(), artifact_dir=tmp_path / "gif-thumbs")
    assert gif_outcome.profile.metadata["frames"] == 5 and len(gif_outcome.cache_artifacts) == 3
    assert "ANIMATED_IMAGE_SAMPLED" in gif_outcome.profile.warnings
    assert_schema(project_root, image_outcome); assert_schema(project_root, gif_outcome)


def test_pa07_real_wav_metadata_and_missing_asr_is_explicit(project_root: Path, tmp_path: Path):
    path = tmp_path / "silence.wav"
    with wave.open(str(path), "wb") as audio:
        audio.setnchannels(1); audio.setsampwidth(2); audio.setframerate(8000)
        audio.writeframes(b"\x00\x00" * 8000)
    outcome = ParserRegistry().parse(path, uid())
    assert outcome.status == "partial" and outcome.profile.metadata["duration_sec"] == 1.0
    assert "ASR_COMPONENT_MISSING" in outcome.profile.warnings
    assert not any(item.kind == "transcript" for item in outcome.profile.evidence)
    assert_schema(project_root, outcome)


def test_pa08_video_without_ffprobe_is_partial_not_fake_success(project_root: Path, tmp_path: Path):
    path = tmp_path / "video.mp4"
    writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 5.0, (160, 120))
    assert writer.isOpened()
    for index in range(15):
        frame = np.full((120, 160, 3), (index * 12, 60, 120), dtype=np.uint8)
        cv2.putText(frame, f"frame {index}", (12, 64), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 1)
        writer.write(frame)
    writer.release()
    assert path.stat().st_size > 1_000
    outcome = ParserRegistry().parse(path, uid())
    assert outcome.status == "partial" and "FFPROBE_COMPONENT_MISSING" in outcome.profile.warnings
    assert outcome.profile.coverage.mode == "metadata_only" and not outcome.profile.evidence
    assert_schema(project_root, outcome)


def test_pa09_stratified_pdf_coverage(project_root: Path, tmp_path: Path):
    path = tmp_path / "many.pdf"; pdf = canvas.Canvas(str(path))
    for page in range(20): pdf.drawString(72, 720, f"Evidence page number {page + 1} with enough deterministic text."); pdf.showPage()
    pdf.save(); outcome = ParserRegistry().parse(path, uid(), "fast")
    assert outcome.profile.coverage.sampled_pages == [1, 7, 14, 20]
    assert outcome.profile.coverage.truncated is True
    assert_schema(project_root, outcome)


def test_pa10_corrupt_signature_and_zip_bomb_limit(project_root: Path, tmp_path: Path):
    fake = tmp_path / "fake.pdf"; fake.write_text("not pdf")
    fake_outcome = ParserRegistry().parse(fake, uid()); assert fake_outcome.status == "failed"; assert "FORMAT_SIGNATURE_MISMATCH" in fake_outcome.profile.warnings
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive: archive.writestr("word/document.xml", b"0" * (2 * 1024 * 1024))
    bomb_outcome = ParserRegistry().parse(bomb, uid()); assert bomb_outcome.status == "failed"; assert "ARCHIVE_EXPANSION_LIMIT" in bomb_outcome.profile.warnings
    huge = tmp_path / "huge.png"; Image.new("1", (8001, 8001), 1).save(huge)
    huge_outcome = ParserRegistry().parse(huge, uid()); assert huge_outcome.status == "failed"; assert "IMAGE_PIXEL_LIMIT" in huge_outcome.profile.warnings
    assert_schema(project_root, fake_outcome); assert_schema(project_root, bomb_outcome); assert_schema(project_root, huge_outcome)


def test_pa11_component_states_and_validated_resource_import(project_root: Path, tmp_path: Path):
    database = Database(tmp_path / "app.sqlite3", project_root / "contracts" / "database.sql"); database.initialize()
    manager = ComponentManager(database, tmp_path / "installed")
    states = {item.component_type: item for item in manager.list()}
    assert states["ocr"].status == "ready"
    assert states["asr"].status in {"missing", "disabled"}
    source = tmp_path / "pack"; source.mkdir(); payload = source / "ffprobe.exe"; payload.write_bytes(b"licensed-test-binary")
    manifest = {
        "component_id": "ffmpeg.test", "component_type": "ffmpeg", "version": "test-1",
        "platform": "windows", "arch": "x86_64", "license_source": "test fixture, CC0",
        "relative_files": [{"path": "ffprobe.exe", "sha256": hashlib.sha256(payload.read_bytes()).hexdigest()}],
        "entry_paths": ["ffprobe.exe"],
    }
    manifest_path = source / "manifest.json"; manifest_path.write_text(json.dumps(manifest, sort_keys=True), "utf-8")
    digest = hashlib.sha256(manifest_path.read_bytes()).hexdigest()
    imported = manager.import_directory(source, "ffmpeg", digest)
    assert imported.status == "ready" and (tmp_path / "installed" / "ffmpeg" / "ffmpeg.test" / "ffprobe.exe").read_bytes() == payload.read_bytes()
    bad = tmp_path / "bad"; bad.mkdir(); (bad / "setup.ps1").write_text("Write-Host unsafe")
    bad_manifest = {**manifest, "component_id": "bad", "relative_files": [{"path": "setup.ps1", "sha256": hashlib.sha256((bad / "setup.ps1").read_bytes()).hexdigest()}], "entry_paths": ["setup.ps1"]}
    (bad / "manifest.json").write_text(json.dumps(bad_manifest, sort_keys=True), "utf-8")
    with pytest.raises(ComponentError, match="COMPONENT_SCRIPT_REJECTED"):
        manager.import_directory(bad, "ffmpeg", hashlib.sha256((bad / "manifest.json").read_bytes()).hexdigest())
    database.close()


def test_pa12_worker_timeout_kills_only_owned_process(tmp_path: Path):
    path = tmp_path / "worker.txt"; path.write_text("timeout fixture")
    pid_file = tmp_path / "worker.pid"
    runner = ParserRunner(tmp_path / "cache", timeout_seconds=0.2)
    runner._command = lambda _job, _output: [sys.executable, "-c", f"import os,time,pathlib; pathlib.Path({str(pid_file)!r}).write_text(str(os.getpid())); time.sleep(60)"]
    started = time.monotonic(); outcome = runner.parse(path, uid()); elapsed = time.monotonic() - started
    assert outcome.status == "failed" and outcome.profile.warnings == ["PARSER_TIMEOUT"] and elapsed < 5
    pid = int(pid_file.read_text()); assert not __import__("psutil").pid_exists(pid)


def test_worker_runner_real_process_and_cache(project_root: Path, tmp_path: Path):
    path = tmp_path / "worker.txt"; path.write_text("worker evidence")
    runner = ParserRunner(tmp_path / "cache", timeout_seconds=10)
    first = runner.parse(path, uid())
    assert first.status == "ready"
    cache_files = list((tmp_path / "cache").glob("*.json")); assert len(cache_files) == 1
    second_id = uid(); second = runner.parse(path, second_id)
    assert second.profile.content_summary == first.profile.content_summary
    assert second.profile.file_id == second_id
