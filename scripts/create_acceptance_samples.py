from __future__ import annotations

import hashlib
import json
import math
import struct
import wave
from pathlib import Path

import cv2
import numpy as np
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "backend" / "tests" / "fixtures" / "evaluation" / "fixtures"
MANIFEST = ROOT / "backend" / "tests" / "fixtures" / "evaluation" / "gold-manifest.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def create_file(kind: str, index: int, path: Path) -> str:
    phrase = f"Guixu synthetic acceptance sample {kind} {index:02d}"
    if kind == "text":
        path.write_text(phrase + "\nPurpose: deterministic technical classification smoke test.\n", encoding="utf-8")
        return "universal.types.text"
    if kind == "image":
        image = Image.new("RGB", (320, 180), ((index * 29) % 255, 90, 130))
        draw = ImageDraw.Draw(image); draw.rectangle((12, 12, 308, 168), outline="white", width=3); draw.text((24, 72), phrase, fill="white")
        image.save(path)
        return "universal.types.image"
    if kind == "pdf":
        document = canvas.Canvas(str(path)); document.drawString(72, 760, phrase); document.drawString(72, 730, "Self-generated; no personal data."); document.save()
        return "universal.types.pdf"
    if kind == "office":
        suffix = path.suffix.lower()
        if suffix == ".docx":
            document = Document(); document.add_heading(phrase, 1); document.add_paragraph("Deterministic Word fixture."); document.save(path)
            return "universal.types.word"
        if suffix == ".pptx":
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = phrase; slide.placeholders[1].text = "Deterministic slides fixture."; deck.save(path)
            return "universal.types.slides"
        book = Workbook(); sheet = book.active; sheet.title = "Acceptance"; sheet.append(["sample", "index"]); sheet.append([phrase, index]); book.save(path)
        return "universal.types.sheet"
    if kind == "audio":
        with wave.open(str(path), "wb") as audio:
            audio.setnchannels(1); audio.setsampwidth(2); audio.setframerate(8000)
            frames = bytearray()
            for sample in range(4000):
                value = int(5000 * math.sin(2 * math.pi * (220 + index) * sample / 8000))
                frames.extend(struct.pack("<h", value))
            audio.writeframes(bytes(frames))
        return "universal.types.audio"
    writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 5.0, (160, 120))
    if not writer.isOpened():
        raise RuntimeError("OpenCV MP4 writer unavailable")
    for frame_index in range(10):
        frame = np.full((120, 160, 3), ((index * 9) % 255, frame_index * 18, 110), dtype=np.uint8)
        cv2.putText(frame, f"{index:02d}-{frame_index}", (24, 68), cv2.FONT_HERSHEY_SIMPLEX, .7, (255, 255, 255), 2)
        writer.write(frame)
    writer.release()
    return "universal.types.video"


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    if any(OUTPUT.iterdir()) or MANIFEST.exists():
        raise SystemExit("acceptance fixture output already exists; refusing to overwrite")
    cases = []
    for kind in ("text", "image", "pdf", "office", "audio", "video"):
        directory = OUTPUT / kind
        directory.mkdir()
        for index in range(1, 21):
            suffix = {
                "text": ".txt", "image": ".png", "pdf": ".pdf", "audio": ".wav", "video": ".mp4",
                "office": (".docx", ".pptx", ".xlsx")[(index - 1) % 3],
            }[kind]
            path = directory / f"{kind}-{index:02d}{suffix}"
            category = create_file(kind, index, path)
            cases.append({
                "id": f"{kind}-{index:02d}", "relative_path": path.relative_to(OUTPUT).as_posix(),
                "sha256": sha256(path), "size_bytes": path.stat().st_size,
                "acceptable_category_ids": [category], "should_abstain": False,
                "evidence_expectation": "locally parsed type or container evidence",
            })
    manifest = {
        "schema_version": 1,
        "license": "CC0-1.0; generated entirely by this project script without third-party or user content",
        "generator": "scripts/create_acceptance_samples.py",
        "purpose": "technical parser/classifier smoke set; not a claim of real-world semantic quality",
        "case_count": len(cases),
        "modalities": {kind: 20 for kind in ("text", "image", "pdf", "office", "audio", "video")},
        "cases": cases,
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (MANIFEST.parent / "LICENSE.txt").write_text(
        "The synthetic fixtures generated by scripts/create_acceptance_samples.py are dedicated to the public domain under CC0-1.0.\nNo third-party or user content is included.\n",
        encoding="utf-8",
    )
    print(f"created {len(cases)} fixtures at {OUTPUT}")


if __name__ == "__main__":
    main()
